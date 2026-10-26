"""
Syllabus ingestion without an LLM: DOCX table extraction + rule-based competency fallback.
"""

import io

from docx import Document

from backend.services.ai.syllabus_parser import _extract_competencies_from_text
from backend.services.ocr.extractor import extract_text_from_docx


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("BALIWAG POLYTECHNIC COLLEGE")
    doc.add_paragraph("Institute of Information Technology and Innovation")
    info = doc.add_table(rows=1, cols=2)
    info.rows[0].cells[0].text, info.rows[0].cells[1].text = "Course Title", "Data Structures and Algorithms"
    doc.add_paragraph("COURSE COMPETENCIES")
    comp = doc.add_table(rows=3, cols=5)
    rows = [
        ("Code", "Competency / Topic", "CLO", "Period", "Learning Outcome"),
        ("DSA-01", "Algorithm Analysis", "CLO1", "Prelim", "Compare growth rates using Big-O."),
        ("DSA-02", "Linked Lists", "CLO2", "Prelim", "Implement singly linked lists."),
    ]
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            comp.rows[r].cells[c].text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extraction_keeps_table_rows_in_order():
    text = extract_text_from_docx(_docx_bytes())
    lines = text.splitlines()
    assert "Course Title | Data Structures and Algorithms" in lines
    assert "DSA-01 | Algorithm Analysis | CLO1 | Prelim | Compare growth rates using Big-O." in lines
    assert lines.index("COURSE COMPETENCIES") < lines.index(next(l for l in lines if l.startswith("DSA-01")))


def test_fallback_uses_coded_table_rows_not_letterhead():
    result = _extract_competencies_from_text(extract_text_from_docx(_docx_bytes()))
    assert result["course_title"] == "Data Structures and Algorithms"
    assert [c["competency_code"] for c in result["competencies"]] == ["DSA-01", "DSA-02"]
    assert result["competencies"][0]["topic_name"] == "Algorithm Analysis"
    assert result["competencies"][0]["learning_outcomes"] == ["Compare growth rates using Big-O."]


def test_fallback_parses_pdf_flattened_rows():
    text = "\n".join([
        "Republic of the Philippines",
        "Course Title Data Structures and Algorithms",
        "DSA-01 Algorithm Analysis and Big-O Notation CLO1 Prelim Determine the growth rate of an algorithm.",
        "DSA-02 Sorting Algorithms CLO4 Midterm Implement merge and quick sort.",
    ])
    result = _extract_competencies_from_text(text)
    assert result["course_title"] == "Data Structures and Algorithms"
    first, second = result["competencies"]
    assert (first["competency_code"], first["topic_name"]) == ("DSA-01", "Algorithm Analysis and Big-O Notation")
    assert first["learning_outcomes"] == ["Determine the growth rate of an algorithm."]
    assert (second["topic_name"], second["learning_outcomes"]) == ("Sorting Algorithms", ["Implement merge and quick sort."])


def test_fallback_uses_clo_lines_when_no_topic_codes():
    text = "\n".join([
        "CLO1 Analyze algorithm complexity using Big-O notation. a, b",
        "CLO2 Implement stacks and queues to solve problems. a, c, i",
    ])
    comps = _extract_competencies_from_text(text)["competencies"]
    assert [c["competency_code"] for c in comps] == ["CLO1", "CLO2"]
    assert comps[0]["topic_name"] == "Analyze algorithm complexity using Big-O notation."


def test_fallback_still_handles_plain_numbered_lists():
    text = "1. Understand the basics of relational databases\n2. Write SQL queries with joins and filters"
    comps = _extract_competencies_from_text(text)["competencies"]
    assert [c["competency_code"] for c in comps] == ["CLO1", "CLO2"]
    assert comps[1]["topic_name"] == "Write SQL queries with joins and filters"
