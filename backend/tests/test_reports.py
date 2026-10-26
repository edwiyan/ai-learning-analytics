"""
Academic Performance & Early Intervention report: role scoping, breakdowns by
section / curriculum, filters, and the PDF / Excel exports.
"""
import io
import uuid

import pytest
from openpyxl import load_workbook

from backend.core.security import create_access_token, hash_password, hash_student_id
from backend.models.assessment import (
    Assessment, AssessmentApprovalStatus, AssessmentItem, AssessmentType, GradingPeriod,
)
from backend.models.course import Course
from backend.models.curriculum import Curriculum
from backend.models.enrollment import StudentEnrollment
from backend.models.program import Program
from backend.models.score import StudentScore
from backend.models.section import Section
from backend.models.syllabus import Competency, Syllabus, SyllabusStatus
from backend.models.user import User, UserRole


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _token(user):
    return create_access_token({"sub": str(user.id), "role": user.role, "jti": str(uuid.uuid4())})


def _user(db, email, role, college=None, **extra):
    u = User(email=email, password_hash=hash_password("Pass@1234"), full_name=email.split("@")[0].title(),
             role=role, is_active=True, college=college, **extra)
    db.add(u)
    db.commit()
    return u


def _course(db, instructor, code, students):
    """students: list of (student_number, pct) — one prelim exam over three competencies."""
    course = Course(title=f"Course {code}", code=code, academic_period="AY 2025-2026 S1", instructor_id=instructor.id)
    db.add(course)
    db.flush()
    syl = Syllabus(course_id=course.id, file_path="x", status=SyllabusStatus.validated)
    db.add(syl)
    db.flush()
    a = Assessment(course_id=course.id, title="Prelim", type=AssessmentType.exam,
                   grading_period=GradingPeriod.prelim, approval_status=AssessmentApprovalStatus.approved)
    db.add(a)
    db.flush()
    for i in range(3):
        comp = Competency(syllabus_id=syl.id, competency_code=f"{code}-0{i + 1}", topic_name=f"{code} topic {i + 1}",
                          learning_outcomes=[])
        db.add(comp)
        db.flush()
        item = AssessmentItem(assessment_id=a.id, question_text=f"Q{i}", mapped_competency_id=comp.id,
                              confidence_score=1.0)
        db.add(item)
        db.flush()
        for number, pct in students:
            db.add(StudentScore(student_id_hash=hash_student_id(number), assessment_id=a.id, item_id=item.id,
                                competency_id=comp.id, grading_period=GradingPeriod.prelim,
                                score=pct, max_score=100))
    db.commit()
    return course


@pytest.fixture
def campus(db, instructor_user, other_instructor_user):
    """ICS course (instructor_user) with two sections in one curriculum; CBA course (other instructor)."""
    other_instructor_user.college = "CBA"
    db.commit()
    prog = Program(program_name="BS Information Technology", degree_type="BS", total_units=150)
    db.add(prog)
    db.flush()
    cur = Curriculum(program_id=prog.id, name="BSIT 2024", effective_year=2024, total_units=150)
    db.add(cur)
    db.flush()
    sec_a = Section(curriculum_id=cur.id, section_name="BSIT 2A", year_level=2, academic_year="2025-2026")
    sec_b = Section(curriculum_id=cur.id, section_name="BSIT 2B", year_level=2, academic_year="2025-2026")
    db.add_all([sec_a, sec_b])
    db.flush()
    for number, name, sec in [("2024-0001", "Ana Weak", sec_a), ("2024-0002", "Ben Strong", sec_a),
                              ("2024-0003", "Cy Middle", sec_b)]:
        _user(db, f"{name.split()[0].lower()}@s.test", UserRole.student, "ICS", student_number=number)
        db.query(User).filter(User.student_number == number).update({"full_name": name})
        db.add(StudentEnrollment(student_id_hash=hash_student_id(number), section_id=sec.id, curriculum_id=cur.id))
    _user(db, "dee@s.test", UserRole.student, "CBA", student_number="2024-0100")
    db.commit()
    ics = _course(db, instructor_user, "ITC 212", [("2024-0001", 50), ("2024-0002", 95), ("2024-0003", 78)])
    cba = _course(db, other_instructor_user, "ACCT 201", [("2024-0100", 60)])
    return {"ics": ics, "cba": cba, "curriculum": cur}


def test_admin_report_covers_everything_with_section_breakdown(client, db, admin_token, campus):
    body = client.get("/reports/academic-performance", headers=_auth(admin_token)).json()
    s = body["summary"]
    assert (s["total"], s["at_risk"], s["on_track"], s["excelling"]) == (4, 2, 1, 1)
    sections = {g["group"]: g for g in body["breakdowns"]["Section"]}
    assert sections["BSIT 2A"]["at_risk"] == 1 and sections["BSIT 2A"]["excelling"] == 1
    assert sections["BSIT 2B"]["on_track"] == 1
    curricula = {g["group"]: g["total"] for g in body["breakdowns"]["Curriculum"]}
    assert curricula["BSIT 2024"] == 3 and curricula["No enrolment record"] == 1
    assert {r["name"] for r in body["at_risk"]} == {"Ana Weak", "Dee"}
    assert body["at_risk"][0]["reason"].startswith("3 competencies below 75%")


def test_instructor_sees_only_own_courses(client, instructor_token, campus):
    body = client.get("/reports/academic-performance", headers=_auth(instructor_token)).json()
    assert body["summary"]["courses"] == 1 and body["summary"]["total"] == 3
    assert body["scope"].startswith("Courses handled by")


def test_dean_sees_only_their_college(client, db, campus):
    dean = _user(db, "dean@test.com", UserRole.dean, "CBA")
    body = client.get("/reports/academic-performance", headers=_auth(_token(dean))).json()
    assert [g["group"] for g in body["breakdowns"]["Course"]][0].startswith("ACCT 201")
    assert body["summary"]["total"] == 1


def test_students_cannot_get_reports(client, student_token, campus):
    assert client.get("/reports/academic-performance", headers=_auth(student_token)).status_code == 403


def test_filters_by_section_and_curriculum(client, admin_token, campus):
    by_section = client.get("/reports/academic-performance", params={"section": "BSIT 2A"},
                            headers=_auth(admin_token)).json()
    assert by_section["summary"]["total"] == 2
    assert ("Section", "BSIT 2A") in [tuple(f) for f in by_section["filters"]]
    by_cur = client.get("/reports/academic-performance", params={"curriculum_id": str(campus["curriculum"].id)},
                        headers=_auth(admin_token)).json()
    assert by_cur["summary"]["total"] == 3
    opts = client.get("/reports/options", headers=_auth(admin_token)).json()
    assert opts["sections"] == ["BSIT 2A", "BSIT 2B", "Unassigned"]
    assert opts["curricula"] == [{"id": str(campus["curriculum"].id), "name": "BSIT 2024"}]


def test_pdf_export(client, admin_token, campus):
    r = client.get("/reports/academic-performance/export", params={"format": "pdf"}, headers=_auth(admin_token))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 5000
    assert "attachment" in r.headers["content-disposition"]


def test_xlsx_export_has_every_view(client, admin_token, campus):
    r = client.get("/reports/academic-performance/export", params={"format": "xlsx"}, headers=_auth(admin_token))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "By College", "By Program", "By Curriculum", "By Year Level", "By Section",
                             "By Course", "All Students", "At-Risk", "Excelling", "Early Warning", "Competency Gaps"]
    at_risk_names = sorted(row[0] for row in wb["At-Risk"].iter_rows(min_row=2, values_only=True))
    assert at_risk_names == ["Ana Weak", "Dee"]


def test_empty_scope_still_renders(client, db):
    instructor = _user(db, "new@test.com", UserRole.instructor, "ICS")
    t = _token(instructor)
    assert client.get("/reports/academic-performance", headers=_auth(t)).json()["summary"]["total"] == 0
    r = client.get("/reports/academic-performance/export", params={"format": "pdf"}, headers=_auth(t))
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
