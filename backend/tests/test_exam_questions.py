"""
Exam questions: generation from the validated syllabus (TOS-driven or per competency),
editing, locking after scores, DOCX export, and the score-sheet -> upload round trip.
"""
import csv
import io

import pytest
from docx import Document

from backend.core.limiter import limiter
from backend.models.assessment import Assessment, AssessmentApprovalStatus, AssessmentType, GradingPeriod
from backend.models.course import Course
from backend.models.syllabus import Competency, Syllabus, SyllabusStatus
from backend.models.tos import TableOfSpecifications
from backend.services.ai import question_generator


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    monkeypatch.setattr(limiter, "enabled", False)
    monkeypatch.setattr(question_generator, "llm_configured", lambda: False)

    def _no_llm(*args, **kwargs):
        raise RuntimeError("LLM unavailable in tests")
    monkeypatch.setattr("backend.routers.assessments.map_question_to_competency", _no_llm)


@pytest.fixture
def setup(db, instructor_user):
    course = Course(title="Data Structures", code="ITC 212", academic_period="AY 2025-2026 S1",
                    instructor_id=instructor_user.id)
    db.add(course)
    db.flush()
    syllabus = Syllabus(course_id=course.id, file_path="x.docx", status=SyllabusStatus.validated)
    db.add(syllabus)
    db.flush()
    comps = {}
    for code, topic, outcomes in [
        ("DSA-01", "Algorithm Analysis", ["Determine the growth rate of an algorithm; compare O(n) and O(n^2)"]),
        ("DSA-02", "Linked Lists", ["Implement singly linked lists"]),
        ("DSA-03", "Stacks and Queues", ["Evaluate postfix expressions using a stack"]),
    ]:
        c = Competency(syllabus_id=syllabus.id, competency_code=code, topic_name=topic, learning_outcomes=outcomes)
        db.add(c)
        db.flush()
        comps[code] = c
    a = Assessment(course_id=course.id, title="Prelim", type=AssessmentType.exam,
                   grading_period=GradingPeriod.prelim, approval_status=AssessmentApprovalStatus.approved)
    db.add(a)
    db.commit()
    return {"assessment": a, "comps": comps, "course": course}


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def test_generate_per_competency_template_fallback(client, instructor_token, setup):
    a, comps = setup["assessment"], setup["comps"]
    res = client.post(f"/assessments/{a.id}/questions/generate", headers=_auth(instructor_token), json={
        "competency_ids": [str(comps["DSA-01"].id), str(comps["DSA-02"].id)], "count_per_competency": 3,
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["source"] == "template" and body["basis"] == "competencies" and body["locked"] is False
    qs = body["questions"]
    assert len(qs) == 6
    assert [q["competency_code"] for q in qs] == ["DSA-01"] * 3 + ["DSA-02"] * 3
    assert len({q["question_text"] for q in qs}) == 6           # no duplicates
    assert all(q["question_type"] == "short_answer" and q["points"] == 5 for q in qs)
    assert "growth rate of an algorithm" in qs[0]["question_text"]


def test_generate_follows_tos_and_replaces(client, db, instructor_token, setup):
    a, comps = setup["assessment"], setup["comps"]
    client.post(f"/assessments/{a.id}/questions/generate", headers=_auth(instructor_token), json={})
    db.add(TableOfSpecifications(assessment_id=a.id, competency_id=comps["DSA-03"].id,
                                 bloom_level="Applying", item_count=2, weight_percent=100))
    db.commit()
    res = client.post(f"/assessments/{a.id}/questions/generate", headers=_auth(instructor_token),
                      json={"question_type": "problem"})
    body = res.json()
    assert body["basis"] == "tos"
    assert [(q["competency_code"], q["bloom_level"], q["question_type"]) for q in body["questions"]] == \
        [("DSA-03", "Applying", "problem")] * 2


def test_llm_output_is_validated(monkeypatch):
    monkeypatch.setattr(question_generator, "llm_configured", lambda: True)
    monkeypatch.setattr(question_generator, "_llm_questions", lambda *a, **k: [
        {"competency_code": "dsa-01", "question_type": "multiple_choice", "question_text": "Big-O of array access?",
         "choices": ["A. O(1)", "B. O(n)", "C. O(log n)", "D. O(n^2)"], "answer_key": "A", "bloom_level": "Remembering"},
        {"competency_code": "XYZ-9", "question_type": "short_answer", "question_text": "Unknown code"},
        {"competency_code": "DSA-01", "question_type": "multiple_choice", "question_text": "No choices", "choices": None},
    ])
    plan = [{"competency": {"competency_code": "DSA-01", "topic_name": "Algo", "learning_outcomes": []},
             "count": 3, "bloom_level": None}]
    qs, source = question_generator.generate_questions(plan)
    assert source == "ai"
    assert len(qs) == 2
    assert qs[0]["choices"] == ["O(1)", "O(n)", "O(log n)", "O(n^2)"] and qs[0]["points"] == 1
    assert qs[1]["question_type"] == "short_answer" and qs[1]["choices"] is None


def test_edit_add_delete_and_validation(client, instructor_token, setup):
    a, comps = setup["assessment"], setup["comps"]
    added = client.post(f"/assessments/{a.id}/questions", headers=_auth(instructor_token), json={
        "question_text": "What is the top after pushing 5, 8, 3 and popping twice?",
        "question_type": "multiple_choice", "choices": ["3", "5", "8", "empty"], "answer_key": "B",
        "points": 1, "competency_id": str(comps["DSA-03"].id),
    })
    assert added.status_code == 201, added.text
    qid = added.json()["id"]

    bad = client.put(f"/assessments/questions/{qid}", headers=_auth(instructor_token), json={"choices": ["only one"]})
    assert bad.status_code == 422

    edited = client.put(f"/assessments/questions/{qid}", headers=_auth(instructor_token),
                        json={"question_type": "short_answer", "points": 3})
    assert edited.status_code == 200
    assert edited.json()["choices"] is None and edited.json()["points"] == 3

    assert client.delete(f"/assessments/questions/{qid}", headers=_auth(instructor_token)).status_code == 204
    assert client.get(f"/assessments/{a.id}/questions", headers=_auth(instructor_token)).json()["questions"] == []


def test_other_instructor_cannot_access(client, other_instructor_token, setup):
    res = client.post(f"/assessments/{setup['assessment'].id}/questions/generate",
                      headers=_auth(other_instructor_token), json={})
    assert res.status_code == 403


def test_docx_export_contains_questions(client, instructor_token, setup):
    a = setup["assessment"]
    qs = client.post(f"/assessments/{a.id}/questions/generate", headers=_auth(instructor_token),
                     json={"count_per_competency": 1}).json()["questions"]
    for answer_key in (False, True):
        res = client.get(f"/assessments/{a.id}/questions/export/docx",
                         params={"answer_key": answer_key}, headers=_auth(instructor_token))
        assert res.status_code == 200
        text = "\n".join(p.text for p in Document(io.BytesIO(res.content)).paragraphs)
        assert all(q["question_text"] in text for q in qs)
        assert ("ANSWER KEY" in text) is answer_key
        assert ("Name:" in "\n".join(c.text for t in Document(io.BytesIO(res.content)).tables
                                     for r in t.rows for c in r.cells)) is (not answer_key)


def test_score_sheet_round_trip_and_lock(client, instructor_token, setup):
    a = setup["assessment"]
    qs = client.post(f"/assessments/{a.id}/questions/generate", headers=_auth(instructor_token),
                     json={"count_per_competency": 1}).json()["questions"]

    sheet = client.get(f"/assessments/{a.id}/questions/export/score-sheet", headers=_auth(instructor_token))
    assert sheet.status_code == 200
    rows = list(csv.DictReader(io.StringIO(sheet.text)))
    assert len(rows) == len(qs)                       # no roster yet: one blank block
    assert all(r["student_id"] == "" and r["competency_code"] for r in rows)

    # Instructor fills in students and scores, then uploads it back
    filled = io.StringIO()
    writer = csv.DictWriter(filled, fieldnames=rows[0].keys())
    writer.writeheader()
    for sid in ("2024-0001", "2024-0002"):
        for r in rows:
            writer.writerow({**r, "student_id": sid, "score": r["max_score"]})
    up = client.post(f"/assessments/{a.id}/upload", headers=_auth(instructor_token),
                     files={"file": ("scores.csv", filled.getvalue().encode(), "text/csv")})
    assert up.status_code == 200, up.text
    assert up.json()["questions"] == len(qs)
    assert up.json()["pre_tagged"] == len(qs)
    assert up.json()["flagged_for_review"] == 0

    listing = client.get(f"/assessments/{a.id}/questions", headers=_auth(instructor_token)).json()
    assert listing["locked"] is True
    locked = client.put(f"/assessments/questions/{qs[0]['id']}", headers=_auth(instructor_token),
                        json={"question_text": "Changed after grading"})
    assert locked.status_code == 409

    # Roster now comes from the uploaded scores (students with accounts only; none here)
    again = client.get(f"/assessments/{a.id}/questions/export/score-sheet", headers=_auth(instructor_token))
    assert again.status_code == 200
