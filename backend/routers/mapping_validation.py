"""Mapping validation (panel comment 1): build a faculty-validated ground truth for the
AI's question-to-competency mapping and score the AI against it.

Roles
  * Managers (create sets, assign people, see results): administrators and Program
    Directors for their college's courses; Deans and the VPAA can view results.
  * Validators: label their course's sampled items blind — the AI's answer and
    confidence are never sent to them, nor the other validator's labels.
  * Adjudicator (usually the Program Director): resolves items the validators disagree on.
"""
import random
import uuid
from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session, selectinload

from backend.core.audit import log_action
from backend.core.database import get_db
from backend.core.deps import get_current_user
from backend.core.notifications import notify
from backend.models.assessment import Assessment, AssessmentItem
from backend.models.course import Course
from backend.models.mapping_validation import (
    MappingValidationAdjudication, MappingValidationAssignment, MappingValidationItem, MappingValidationLabel,
    MappingValidationSet,
)
from backend.models.syllabus import Competency, Syllabus, SyllabusStatus
from backend.models.user import User, UserRole
from backend.services.institution_settings import get_settings
from backend.services.mapping_validation import (
    DEFAULT_SAMPLE_SIZE, NONE, compute_metrics, draw_sample, eligible_items, ground_truth, label_key,
)

router = APIRouter(prefix="/mapping-validation", tags=["mapping-validation"])

MANAGER_ROLES = {UserRole.admin, UserRole.program_director}
VIEWER_ROLES = {UserRole.admin, UserRole.program_director, UserRole.dean, UserRole.vp_academic_affairs}
ACCURACY_TARGET = 85.0


class SetCreate(BaseModel):
    name: str = Field(min_length=3, max_length=200)
    course_ids: list[uuid.UUID] = Field(min_length=1)
    sample_size: int = Field(default=DEFAULT_SAMPLE_SIZE, ge=10, le=5000)
    seed: int | None = None


class AssignmentIn(BaseModel):
    validator_a_id: uuid.UUID | None = None
    validator_b_id: uuid.UUID | None = None
    adjudicator_id: uuid.UUID | None = None


class LabelIn(BaseModel):
    competency_id: uuid.UUID | None = None
    none_fits: bool = False
    note: str | None = Field(default=None, max_length=1000)


# ── helpers ──────────────────────────────────────────────────────────────────

def _course_college(course: Course, db: Session) -> str | None:
    inst = db.get(User, course.instructor_id)
    return inst.college if inst else None


def _courses_in_scope(db: Session, user: User) -> list[Course]:
    q = db.query(Course)
    if user.role in (UserRole.program_director, UserRole.dean):
        q = q.join(User, Course.instructor_id == User.id).filter(User.college == user.college)
    elif user.role not in (UserRole.admin, UserRole.vp_academic_affairs):
        return []
    return q.order_by(Course.code).all()


def _get_set(set_id: uuid.UUID, db: Session) -> MappingValidationSet:
    s = (db.query(MappingValidationSet)
         .options(selectinload(MappingValidationSet.assignments),
                  selectinload(MappingValidationSet.items).selectinload(MappingValidationItem.labels),
                  selectinload(MappingValidationSet.items).selectinload(MappingValidationItem.adjudication))
         .filter(MappingValidationSet.id == set_id).first())
    if not s:
        raise HTTPException(status_code=404, detail="Validation set not found")
    return s


def _can_view(s: MappingValidationSet, user: User, db: Session) -> bool:
    if user.role in (UserRole.admin, UserRole.vp_academic_affairs):
        return True
    if user.role in (UserRole.program_director, UserRole.dean):
        return any(_course_college(db.get(Course, a.course_id), db) == user.college for a in s.assignments)
    return False


def _can_manage(s: MappingValidationSet, user: User, db: Session) -> bool:
    if user.role == UserRole.admin:
        return True
    if user.role == UserRole.program_director:
        return all(_course_college(db.get(Course, a.course_id), db) == user.college for a in s.assignments)
    return False


def _competencies(course_id, db: Session) -> list[Competency]:
    return (db.query(Competency).join(Syllabus, Competency.syllabus_id == Syllabus.id)
            .filter(Syllabus.course_id == course_id, Syllabus.status == SyllabusStatus.validated)
            .order_by(Competency.competency_code).all())


def _comp_out(c: Competency) -> dict:
    return {"id": str(c.id), "competency_code": c.competency_code, "topic_name": c.topic_name,
            "learning_outcomes": c.learning_outcomes or []}


def _label_out(key: str | None, names: dict) -> dict | None:
    if key is None:
        return None
    if key == NONE:
        return {"competency_id": None, "none_fits": True, "label": "No matching competency"}
    code, topic = names.get(key, ("?", "Unknown competency"))
    return {"competency_id": key, "none_fits": False, "label": f"{code} · {topic}"}


def _user_name(uid, db):
    u = db.get(User, uid) if uid else None
    return u.full_name if u else None


def _summary(s: MappingValidationSet, db: Session, user: User, with_metrics: bool) -> dict:
    by_course = {str(a.course_id): a for a in s.assignments}
    out = {
        "id": str(s.id), "name": s.name, "sample_size": s.sample_size, "seed": s.seed,
        "review_threshold": s.review_threshold, "created_at": s.created_at.isoformat() if s.created_at else None,
        "created_by": _user_name(s.created_by, db), "items": len(s.items),
        "can_manage": _can_manage(s, user, db), "can_view_results": _can_view(s, user, db),
        "courses": [],
    }
    for course_id, a in by_course.items():
        course = db.get(Course, a.course_id)
        items = [i for i in s.items if str(i.course_id) == course_id]
        done = lambda vid: sum(1 for i in items if any(str(l.validator_id) == str(vid) for l in i.labels)) if vid else 0
        disagree = [i for i in items if ground_truth(i, a.validator_a_id, a.validator_b_id) is None
                    and len({label_key(l.competency_id, l.none_fits) for l in i.labels
                             if str(l.validator_id) in (str(a.validator_a_id), str(a.validator_b_id))}) == 2]
        out["courses"].append({
            "course_id": course_id,
            "course": f"{course.code} — {course.title}" if course else course_id,
            "items": len(items),
            "validator_a": {"id": str(a.validator_a_id) if a.validator_a_id else None,
                            "name": _user_name(a.validator_a_id, db), "labeled": done(a.validator_a_id)},
            "validator_b": {"id": str(a.validator_b_id) if a.validator_b_id else None,
                            "name": _user_name(a.validator_b_id, db), "labeled": done(a.validator_b_id)},
            "adjudicator": {"id": str(a.adjudicator_id) if a.adjudicator_id else None,
                            "name": _user_name(a.adjudicator_id, db)},
            "awaiting_adjudication": len(disagree),
            "my_role": ("validator" if str(user.id) in (str(a.validator_a_id), str(a.validator_b_id)) else
                        "adjudicator" if str(user.id) == str(a.adjudicator_id) else None),
        })
    if with_metrics:
        names = {str(c.id): (c.competency_code, c.topic_name)
                 for cid in by_course for c in _competencies(cid, db)}
        out["metrics"] = compute_metrics(s.items, by_course, names, ACCURACY_TARGET)
    return out


# ── set management ───────────────────────────────────────────────────────────

@router.get("/options")
def creation_options(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if user.role not in MANAGER_ROLES:
        raise HTTPException(status_code=403, detail="Only administrators and Program Directors create validation sets")
    courses = _courses_in_scope(db, user)
    counts = {}
    if courses:
        for it in eligible_items(db, [c.id for c in courses]):
            counts[str(it.assessment.course_id)] = counts.get(str(it.assessment.course_id), 0) + 1
    staff_q = db.query(User).filter(User.is_active == True, User.role.in_(  # noqa: E712
        [UserRole.instructor, UserRole.program_director, UserRole.dean]))
    if user.role == UserRole.program_director:
        staff_q = staff_q.filter(User.college == user.college)
    staff = staff_q.order_by(User.full_name).all()
    return {
        "default_sample_size": DEFAULT_SAMPLE_SIZE,
        "courses": [{"id": str(c.id), "label": f"{c.code} — {c.title}", "instructor_id": str(c.instructor_id),
                     "eligible_items": counts.get(str(c.id), 0)} for c in courses],
        "validators": [{"id": str(u.id), "name": u.full_name, "role": u.role.value, "college": u.college}
                       for u in staff if u.role in (UserRole.instructor, UserRole.program_director)],
        "adjudicators": [{"id": str(u.id), "name": u.full_name, "role": u.role.value, "college": u.college}
                         for u in staff if u.role in (UserRole.program_director, UserRole.dean)],
    }


@router.post("/", status_code=status.HTTP_201_CREATED)
def create_set(payload: SetCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if user.role not in MANAGER_ROLES:
        raise HTTPException(status_code=403, detail="Only administrators and Program Directors create validation sets")
    allowed = {str(c.id): c for c in _courses_in_scope(db, user)}
    courses = []
    for cid in dict.fromkeys(str(c) for c in payload.course_ids):
        if cid not in allowed:
            raise HTTPException(status_code=403, detail="A selected course is outside your scope")
        courses.append(allowed[cid])
    items = eligible_items(db, [c.id for c in courses])
    if not items:
        raise HTTPException(status_code=422, detail="The selected courses have no AI-mapped assessment items yet")
    threshold = get_settings(db).mapping_review_threshold
    seed = payload.seed if payload.seed is not None else random.SystemRandom().randint(1, 2**31 - 1)
    course_of = {str(i.id): str(i.assessment.course_id) for i in items}
    sample = draw_sample(items, course_of, payload.sample_size, seed, threshold)

    s = MappingValidationSet(name=payload.name.strip(), sample_size=payload.sample_size, seed=seed,
                             review_threshold=threshold, created_by=user.id)
    db.add(s)
    db.flush()
    for pos, (it, course_id, band) in enumerate(sample):
        db.add(MappingValidationItem(
            set_id=s.id, course_id=uuid.UUID(course_id), assessment_item_id=it.id, question_text=it.question_text,
            ai_competency_id=it.ai_suggested_competency_id, ai_confidence=it.ai_confidence,
            confidence_band=band, position=pos))
    for c in courses:
        if any(cid == str(c.id) for _, cid, _ in sample):
            db.add(MappingValidationAssignment(
                set_id=s.id, course_id=c.id, validator_a_id=c.instructor_id,
                adjudicator_id=user.id if user.role == UserRole.program_director else None))
    db.commit()
    log_action(db, str(user.id), "create_mapping_validation_set", target_table="mapping_validation_sets",
               target_id=s.id, extra={"items": len(sample), "courses": len(courses), "seed": seed})
    return _summary(_get_set(s.id, db), db, user, with_metrics=False)


@router.get("/")
def list_sets(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    sets = db.query(MappingValidationSet).order_by(MappingValidationSet.created_at.desc()).all()
    out = []
    for s in sets:
        s = _get_set(s.id, db)
        involved = any(str(user.id) in (str(a.validator_a_id), str(a.validator_b_id), str(a.adjudicator_id))
                       for a in s.assignments)
        if involved or _can_view(s, user, db):
            out.append(_summary(s, db, user, with_metrics=False))
    return out


@router.get("/{set_id}")
def get_set(set_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = _get_set(set_id, db)
    viewer = _can_view(s, user, db)
    involved = any(str(user.id) in (str(a.validator_a_id), str(a.validator_b_id), str(a.adjudicator_id))
                   for a in s.assignments)
    if not (viewer or involved):
        raise HTTPException(status_code=403, detail="Not part of this validation set")
    return _summary(s, db, user, with_metrics=viewer)


@router.put("/{set_id}/assignments/{course_id}")
def assign(set_id: uuid.UUID, course_id: uuid.UUID, payload: AssignmentIn,
           db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = _get_set(set_id, db)
    if not _can_manage(s, user, db):
        raise HTTPException(status_code=403, detail="Only the set's managers can assign people")
    a = next((a for a in s.assignments if str(a.course_id) == str(course_id)), None)
    if not a:
        raise HTTPException(status_code=404, detail="Course is not part of this set")
    if payload.validator_a_id and payload.validator_a_id == payload.validator_b_id:
        raise HTTPException(status_code=422, detail="The two validators must be different people")
    if payload.adjudicator_id and payload.adjudicator_id in (payload.validator_a_id, payload.validator_b_id):
        raise HTTPException(status_code=422, detail="The adjudicator cannot also be a validator")
    roles = {"validator_a_id": (UserRole.instructor, UserRole.program_director),
             "validator_b_id": (UserRole.instructor, UserRole.program_director),
             "adjudicator_id": (UserRole.program_director, UserRole.dean, UserRole.admin)}
    newly = []
    for field, allowed in roles.items():
        uid = getattr(payload, field)
        if uid is not None:
            u = db.get(User, uid)
            if not u or u.role not in allowed or not u.is_active:
                raise HTTPException(status_code=422, detail=f"{field.replace('_id', '').replace('_', ' ')}: not an eligible user")
        previous = getattr(a, field)
        labeled = previous and db.query(MappingValidationLabel).join(MappingValidationItem).filter(
            MappingValidationItem.set_id == s.id, MappingValidationItem.course_id == a.course_id,
            MappingValidationLabel.validator_id == previous).count()
        if field != "adjudicator_id" and labeled and str(previous) != str(uid):
            raise HTTPException(status_code=409, detail="A validator who has already labeled items cannot be replaced")
        if str(previous) != str(uid) and uid is not None:
            newly.append((uid, field))
        setattr(a, field, uid)
    course = db.get(Course, a.course_id)
    for uid, field in newly:
        role = "adjudicate disagreements" if field == "adjudicator_id" else "label items blind"
        notify(db, str(uid), "validation_assigned", "Mapping validation assignment",
               f"You were asked to {role} for {course.code} in the validation set \"{s.name}\". "
               "Open Mapping Validation to start.", related_id=str(s.id))
    db.commit()
    log_action(db, str(user.id), "assign_mapping_validators", target_table="mapping_validation_assignments",
               target_id=a.id, extra={k: str(v) if v else None for k, v in payload.model_dump().items()})
    return _summary(_get_set(set_id, db), db, user, with_metrics=_can_view(s, user, db))


@router.delete("/{set_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_set(set_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = _get_set(set_id, db)
    if not _can_manage(s, user, db):
        raise HTTPException(status_code=403, detail="Only the set's managers can delete it")
    db.delete(s)
    db.commit()
    log_action(db, str(user.id), "delete_mapping_validation_set", target_table="mapping_validation_sets",
               target_id=set_id)


# ── blind labeling ───────────────────────────────────────────────────────────

def _assignment_for_validator(s, course_id, user):
    a = next((a for a in s.assignments if str(a.course_id) == str(course_id)), None)
    if not a or str(user.id) not in (str(a.validator_a_id), str(a.validator_b_id)):
        raise HTTPException(status_code=403, detail="You are not a validator for this course")
    return a


@router.get("/{set_id}/label/{course_id}")
def labeling_queue(set_id: uuid.UUID, course_id: uuid.UUID, db: Session = Depends(get_db),
                   user: User = Depends(get_current_user)):
    """The validator's items for one course. Deliberately excludes the AI's answer,
    its confidence and the other validator's labels."""
    s = _get_set(set_id, db)
    _assignment_for_validator(s, course_id, user)
    items = sorted((i for i in s.items if str(i.course_id) == str(course_id)), key=lambda i: i.position)
    mine = {}
    for i in items:
        own = next((l for l in i.labels if str(l.validator_id) == str(user.id)), None)
        mine[str(i.id)] = (label_key(own.competency_id, own.none_fits) if own else None,
                           i.adjudication is not None)
    course = db.get(Course, course_id)
    return {
        "set_name": s.name,
        "course": f"{course.code} — {course.title}" if course else str(course_id),
        "competencies": [_comp_out(c) for c in _competencies(course_id, db)],
        "items": [{"id": str(i.id), "question_text": i.question_text,
                   "my_label": mine[str(i.id)][0], "locked": mine[str(i.id)][1]} for i in items],
        "done": sum(1 for v in mine.values() if v[0] is not None),
        "total": len(items),
    }


@router.put("/{set_id}/label/items/{item_id}")
def save_label(set_id: uuid.UUID, item_id: uuid.UUID, payload: LabelIn, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)):
    s = _get_set(set_id, db)
    item = next((i for i in s.items if str(i.id) == str(item_id)), None)
    if not item:
        raise HTTPException(status_code=404, detail="Item not found in this set")
    _assignment_for_validator(s, item.course_id, user)
    if item.adjudication is not None:
        raise HTTPException(status_code=409, detail="This item has been adjudicated; labels are final")
    if not payload.none_fits:
        if payload.competency_id is None:
            raise HTTPException(status_code=422, detail="Choose a competency or 'no matching competency'")
        if str(payload.competency_id) not in {str(c.id) for c in _competencies(item.course_id, db)}:
            raise HTTPException(status_code=422, detail="Competency is not in this course's endorsed syllabus")
    label = next((l for l in item.labels if str(l.validator_id) == str(user.id)), None)
    if label is None:
        label = MappingValidationLabel(item_id=item.id, validator_id=user.id)
        db.add(label)
    label.competency_id = None if payload.none_fits else payload.competency_id
    label.none_fits = payload.none_fits
    label.labeled_at = datetime.now(timezone.utc)
    db.commit()
    return {"item_id": str(item.id), "label": label_key(label.competency_id, label.none_fits)}


# ── adjudication ─────────────────────────────────────────────────────────────

@router.get("/{set_id}/adjudication")
def adjudication_queue(set_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Items where both validators answered differently, for courses this user adjudicates.
    The AI's answer is still hidden so the decision is about the question alone."""
    s = _get_set(set_id, db)
    courses = [a for a in s.assignments
               if str(a.adjudicator_id) == str(user.id) or user.role == UserRole.admin]
    if not courses:
        raise HTTPException(status_code=403, detail="You are not an adjudicator in this set")
    out = []
    for a in courses:
        comps = _competencies(a.course_id, db)
        names = {str(c.id): (c.competency_code, c.topic_name) for c in comps}
        course = db.get(Course, a.course_id)
        rows = []
        for i in sorted((i for i in s.items if str(i.course_id) == str(a.course_id)), key=lambda i: i.position):
            labels = {str(l.validator_id): label_key(l.competency_id, l.none_fits) for l in i.labels}
            la, lb = labels.get(str(a.validator_a_id)), labels.get(str(a.validator_b_id))
            if la is None or lb is None or la == lb:
                continue
            adj = i.adjudication
            rows.append({"id": str(i.id), "question_text": i.question_text,
                         "label_a": _label_out(la, names), "label_b": _label_out(lb, names),
                         "decision": _label_out(label_key(adj.competency_id, adj.none_fits), names) if adj else None,
                         "note": adj.note if adj else None})
        out.append({"course_id": str(a.course_id), "course": f"{course.code} — {course.title}" if course else "",
                    "competencies": [_comp_out(c) for c in comps], "items": rows})
    return {"set_name": s.name, "courses": out}


@router.put("/{set_id}/adjudication/items/{item_id}")
def adjudicate(set_id: uuid.UUID, item_id: uuid.UUID, payload: LabelIn, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)):
    s = _get_set(set_id, db)
    item = next((i for i in s.items if str(i.id) == str(item_id)), None)
    if not item:
        raise HTTPException(status_code=404, detail="Item not found in this set")
    a = next((a for a in s.assignments if str(a.course_id) == str(item.course_id)), None)
    if not a or (str(a.adjudicator_id) != str(user.id) and user.role != UserRole.admin):
        raise HTTPException(status_code=403, detail="You are not the adjudicator for this course")
    labels = {str(l.validator_id): label_key(l.competency_id, l.none_fits) for l in item.labels}
    la, lb = labels.get(str(a.validator_a_id)), labels.get(str(a.validator_b_id))
    if la is None or lb is None or la == lb:
        raise HTTPException(status_code=409, detail="Only items the two validators disagree on are adjudicated")
    if not payload.none_fits:
        if payload.competency_id is None or str(payload.competency_id) not in {str(c.id) for c in _competencies(item.course_id, db)}:
            raise HTTPException(status_code=422, detail="Choose a competency from the endorsed syllabus or 'no matching competency'")
    adj = item.adjudication or MappingValidationAdjudication(item_id=item.id)
    if item.adjudication is None:
        db.add(adj)
    adj.adjudicator_id = user.id
    adj.competency_id = None if payload.none_fits else payload.competency_id
    adj.none_fits = payload.none_fits
    adj.note = (payload.note or "").strip() or None
    adj.decided_at = datetime.now(timezone.utc)
    db.commit()
    log_action(db, str(user.id), "adjudicate_mapping_item", target_table="mapping_validation_items",
               target_id=item.id, extra={"decision": label_key(adj.competency_id, adj.none_fits)})
    return {"item_id": str(item.id), "decision": label_key(adj.competency_id, adj.none_fits)}


# ── results ──────────────────────────────────────────────────────────────────

@router.get("/{set_id}/export")
def export_results(set_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    s = _get_set(set_id, db)
    if not _can_view(s, user, db):
        raise HTTPException(status_code=403, detail="Not allowed to view these results")
    summary = _summary(s, db, user, with_metrics=True)
    m = summary["metrics"]
    by_course = {str(a.course_id): a for a in s.assignments}
    names, codes = {}, {}
    for cid in by_course:
        for c in _competencies(cid, db):
            names[str(c.id)] = (c.competency_code, c.topic_name)
        course = db.get(Course, cid)
        codes[cid] = course.code if course else cid

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1E1B4B"), Font(bold=True, color="FFFFFF")
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Validation set", s.name), ("Items sampled", m["items"]), ("Sampling seed", s.seed),
        ("Confidence band boundary", s.review_threshold), ("Items with ground truth", m["ground_truth_items"]),
        ("AI correct", m["correct"]), ("Mapping accuracy %", m["accuracy_pct"]),
        ("95% CI low %", m["ci_low"]), ("95% CI high %", m["ci_high"]), ("Target %", m["target_pct"]),
        ("Cohen's kappa (validators, before adjudication)", m["kappa"]),
        ("Validator agreement %", m["agreement_pct"]), ("Disagreements", m["disagreements"]),
        ("Adjudicated", m["adjudicated"]),
        ("High-confidence accuracy %", m["by_band"]["high"]["accuracy_pct"]),
        ("High-confidence items", m["by_band"]["high"]["items"]),
        ("Low-confidence accuracy %", m["by_band"]["low"]["accuracy_pct"]),
        ("Low-confidence items", m["by_band"]["low"]["items"]),
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 30

    items_ws = wb.create_sheet("Items")
    header = ["Course", "Question", "AI competency", "AI confidence", "Band", "Validator A", "Validator B",
              "Adjudication", "Ground truth", "AI correct"]
    items_ws.append(header)
    for cell in items_ws[1]:
        cell.fill, cell.font = head_fill, head_font
    lab = lambda key: "—" if key is None else ("No matching competency" if key == NONE
                                               else " · ".join(names.get(key, ("?", "?"))))
    for i in sorted(s.items, key=lambda i: (codes.get(str(i.course_id), ""), i.position)):
        a = by_course.get(str(i.course_id))
        labels = {str(l.validator_id): label_key(l.competency_id, l.none_fits) for l in i.labels}
        gt = ground_truth(i, a.validator_a_id, a.validator_b_id) if a else None
        ai = label_key(i.ai_competency_id, i.ai_competency_id is None)
        items_ws.append([
            codes.get(str(i.course_id)), i.question_text, lab(ai), i.ai_confidence, i.confidence_band,
            lab(labels.get(str(a.validator_a_id))) if a else "", lab(labels.get(str(a.validator_b_id))) if a else "",
            lab(label_key(i.adjudication.competency_id, i.adjudication.none_fits)) if i.adjudication else "",
            lab(gt), "" if gt is None else ("Yes" if ai == gt else "No"),
        ])
    for col, w in zip("ABCDEFGHIJ", (10, 60, 34, 12, 8, 34, 34, 34, 34, 10)):
        items_ws.column_dimensions[col].width = w
    items_ws.freeze_panes = "A2"

    comp_ws = wb.create_sheet("By Competency")
    comp_ws.append(["Code", "Competency", "Items", "Correct", "Accuracy %"])
    for cell in comp_ws[1]:
        cell.fill, cell.font = head_fill, head_font
    for r in m["by_competency"]:
        comp_ws.append([r["competency_code"], r["topic_name"], r["items"], r["correct"], r["accuracy_pct"]])
    comp_ws.column_dimensions["B"].width = 44

    buf = BytesIO()
    wb.save(buf)
    log_action(db, str(user.id), "export_mapping_validation", target_table="mapping_validation_sets", target_id=s.id)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="AILAS_Mapping_Validation_{s.seed}.xlsx"'})
