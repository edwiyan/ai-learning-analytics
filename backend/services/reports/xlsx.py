"""Render the Academic Performance & Early Intervention report as an Excel workbook
(one sheet per view, filterable, for further analysis)."""
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PH_TZ = ZoneInfo("Asia/Manila")
NAVY = "1E1B4B"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
THIN = Side(style="thin", color="E5E7EB")
RISK_FILLS = {
    "At-Risk": (PatternFill("solid", fgColor="FEE2E2"), Font(color="B91C1C", bold=True)),
    "On-Track": (PatternFill("solid", fgColor="FEF3C7"), Font(color="B45309")),
    "Excelling": (PatternFill("solid", fgColor="D1FAE5"), Font(color="047857")),
}


def _sheet(wb, title, header, rows, widths, risk_col=None, pct_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.column - 1 in pct_cols and isinstance(c.value, (int, float)):
                c.number_format = '0.0"%"'
        if risk_col is not None:
            cell = r[risk_col]
            if cell.value in RISK_FILLS:
                cell.fill, cell.font = RISK_FILLS[cell.value]
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def render_xlsx(data: dict) -> bytes:
    wb = Workbook()
    s = data["summary"]
    generated = data["generated_at"].astimezone(PH_TZ)

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = data["title"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = data["institution"]
    ws["A2"].font = Font(color="6B7280")
    meta = [("Prepared for", data["prepared_for"]), ("Scope", data["scope"]),
            ("Academic period", ", ".join(data["academic_periods"]) or "—"),
            ("Generated", f"{generated:%d %B %Y %I:%M %p} PHT")] + [(f"Filter: {k}", v) for k, v in data["filters"]]
    r = 4
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Measure").font = HEADER_FONT
    ws.cell(row=r, column=2, value="Value").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = ws.cell(row=r, column=2).fill = HEADER_FILL
    measures = [
        ("Students assessed", s["students"]), ("Courses", s["courses"]), ("Course enrolments", s["total"]),
        ("At-Risk enrolments", s["at_risk"]), ("At-Risk %", s["at_risk_pct"]),
        ("Distinct students At-Risk", s["students_flagged"]),
        ("On-Track enrolments", s["on_track"]), ("On-Track %", s["on_track_pct"]),
        ("Excelling enrolments", s["excelling"]), ("Excelling %", s["excelling_pct"]),
        ("Average score %", s["avg_score"]),
        ("At-Risk with an intervention", data["interventions"]["with_intervention"]),
        ("At-Risk without an intervention", data["interventions"]["without_intervention"]),
    ]
    for k, v in measures:
        r += 1
        ws.cell(row=r, column=1, value=k)
        c = ws.cell(row=r, column=2, value=v)
        if "%" in k and isinstance(v, (int, float)):
            c.number_format = '0.0"%"'
    ws.cell(row=r + 2, column=1,
            value="Confidential — contains student information protected under RA 10173 (Data Privacy Act of 2012).").font = \
        Font(italic=True, color="B91C1C")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    # Breakdown sheets, each with a stacked bar chart
    for key, label in (("College", "By College"), ("Program", "By Program"), ("Curriculum", "By Curriculum"),
                       ("Year Level", "By Year Level"), ("Section", "By Section"), ("Course", "By Course")):
        groups = data["breakdowns"][key]
        sh = _sheet(wb, label, [key, "Enrolments", "At-Risk", "On-Track", "Excelling", "At-Risk %", "Excelling %",
                                "Avg score %"],
                    [[g["group"], g["total"], g["at_risk"], g["on_track"], g["excelling"], g["at_risk_pct"],
                      g["excelling_pct"], g["avg_score"]] for g in groups],
                    [46, 12, 10, 10, 11, 11, 12, 12], pct_cols=(5, 6, 7))
        if groups:
            chart = BarChart()
            chart.type, chart.grouping, chart.overlap = "bar", "stacked", 100
            chart.title = f"Risk distribution {label.lower()}"
            chart.y_axis.title = "Enrolments"
            n = len(groups) + 1
            chart.add_data(Reference(sh, min_col=3, max_col=5, min_row=1, max_row=n), titles_from_data=True)
            chart.set_categories(Reference(sh, min_col=1, min_row=2, max_row=n))
            for series, color in zip(chart.series, ("DC2626", "D97706", "059669")):
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color
            chart.height, chart.width = max(6, 1 + len(groups) * 0.8), 18
            sh.add_chart(chart, "J2")

    rows = data["rows"]
    student_header = ["Student", "Student no.", "College", "Program", "Curriculum", "Year level", "Section",
                      "Course", "Course title", "Instructor", "Academic period", "Avg score %", "Learning gaps",
                      "Grade decline %", "Risk level", "Why", "Weakest competencies", "Latest intervention"]
    student_widths = [22, 12, 9, 30, 22, 10, 11, 10, 30, 20, 18, 11, 9, 11, 11, 30, 50, 14]

    def student_row(r):
        return [r["name"], r["student_number"], r["college"], r["program"], r["curriculum"], r["year_level"],
                r["section"], r["course"], r["course_title"], r["instructor"], r["academic_period"], r["avg_score"],
                r["gap_count"], r["grade_decline_pct"], r["risk_level"], r["reason"], "; ".join(r["weak_topics"]),
                r["intervention"] or ""]

    _sheet(wb, "All Students", student_header, [student_row(r) for r in rows], student_widths,
           risk_col=14, pct_cols=(11, 13))
    _sheet(wb, "At-Risk", student_header, [student_row(r) for r in rows if r["risk_level"] == "At-Risk"],
           student_widths, risk_col=14, pct_cols=(11, 13))
    _sheet(wb, "Excelling", student_header, [student_row(r) for r in rows if r["risk_level"] == "Excelling"],
           student_widths, risk_col=14, pct_cols=(11, 13))
    _sheet(wb, "Competency Gaps", ["Course", "Code", "Competency", "Students", "Avg mastery %", "Below 75%",
                                   "% below 75%"],
           [[g["course"], g["code"], g["topic"], g["students"], g["avg_mastery"], g["below"], g["below_pct"]]
            for g in data["competency_gaps"]], [11, 10, 44, 10, 14, 11, 12], pct_cols=(4, 6))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
