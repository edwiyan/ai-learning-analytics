"""Printable exam / answer key (DOCX) and pre-tagged score sheet (CSV) for an assessment's questions."""
import csv
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

INSTITUTION = "Baliwag Polytechnic College"
PERIOD_LABELS = {"prelim": "Preliminary", "midterm": "Midterm", "semi_final": "Semi-Final", "final": "Final"}
TYPE_LABELS = {"quiz": "Quiz", "exam": "Examination", "assignment": "Assignment"}
PARTS = [
    ("multiple_choice", "Multiple Choice", "Choose the letter of the best answer. Write it on the space provided."),
    ("short_answer", "Short Answer", "Answer each question briefly but completely."),
    ("problem", "Problem Solving", "Solve each problem. Show your complete solution."),
]
SCORE_SHEET_COLUMNS = ["student_id", "question_text", "score", "max_score", "submission_date", "competency_code"]


def _enum_value(v) -> str:
    return getattr(v, "value", v) or ""


def _fmt_points(p: float) -> str:
    return f"{p:g} pt" + ("" if p == 1 else "s")


def build_exam_docx(course, assessment, questions: list, with_answers: bool = False) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    def centered(text, size=11, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)

    period = PERIOD_LABELS.get(_enum_value(assessment.grading_period), "")
    kind = TYPE_LABELS.get(_enum_value(assessment.type), "Assessment")
    total = sum(q.points for q in questions)

    centered(INSTITUTION.upper(), 13, bold=True)
    centered(f"{course.code} — {course.title}", 11)
    centered(course.academic_period, 10)
    doc.add_paragraph()
    centered(f"{assessment.title} · {period} {kind}".strip(), 13, bold=True)
    if with_answers:
        centered("ANSWER KEY — FOR INSTRUCTOR USE ONLY", 11, bold=True)
    doc.add_paragraph()

    if not with_answers:
        info = doc.add_table(rows=2, cols=2)
        info.rows[0].cells[0].text = "Name: ______________________________"
        info.rows[0].cells[1].text = "Score: ______ / " + f"{total:g}"
        info.rows[1].cells[0].text = "Student No.: _______________________"
        info.rows[1].cells[1].text = "Section: ___________  Date: __________"
        doc.add_paragraph()
        doc.add_paragraph(
            "General Instructions: Read each item carefully. Write your answers legibly. "
            "Any form of cheating will result in a score of zero for this examination."
        ).runs[0].italic = True

    number = 0
    roman = ["I", "II", "III"]
    part_no = 0
    for qtype, label, instruction in PARTS:
        part_qs = [q for q in questions if q.question_type == qtype]
        if not part_qs:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(f"Part {roman[part_no]}. {label} ({_fmt_points(sum(q.points for q in part_qs))})")
        r.bold = True
        part_no += 1
        doc.add_paragraph(instruction).runs[0].italic = True

        for q in part_qs:
            number += 1
            prefix = "_____ " if qtype == "multiple_choice" and not with_answers else ""
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(f"{prefix}{number}. ").bold = True
            p.add_run(q.question_text)
            p.add_run(f"  ({_fmt_points(q.points)})").italic = True

            if qtype == "multiple_choice":
                for i, choice in enumerate(q.choices or []):
                    c = doc.add_paragraph(f"{chr(65 + i)}. {choice}")
                    c.paragraph_format.left_indent = Cm(1.2)
                    c.paragraph_format.space_after = Pt(0)
            elif not with_answers:
                lines = 6 if qtype == "problem" else 3
                for _ in range(lines):
                    blank = doc.add_paragraph("_" * 76)
                    blank.paragraph_format.space_after = Pt(0)

            if with_answers:
                code = q.competency.competency_code if q.competency else "—"
                a = doc.add_paragraph()
                a.paragraph_format.left_indent = Cm(1.2)
                a.add_run("Answer: ").bold = True
                a.add_run(q.answer_key or "(no answer key yet)")
                a.add_run(f"   [{code}{' · ' + q.bloom_level if q.bloom_level else ''}]").italic = True

    if with_answers:
        doc.add_paragraph()
        doc.add_paragraph().add_run("Coverage by Competency").bold = True
        coverage = {}
        for q in questions:
            key = (q.competency.competency_code, q.competency.topic_name) if q.competency else ("—", "Untagged")
            items, pts = coverage.get(key, (0, 0.0))
            coverage[key] = (items + 1, pts + q.points)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for cell, h in zip(t.rows[0].cells, ["Code", "Competency", "Items", "Points"]):
            cell.text = h
        for (code, topic), (items, pts) in sorted(coverage.items()):
            row = t.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = code, topic, str(items), f"{pts:g}"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_score_sheet_csv(questions: list, student_numbers: list[str], submission_date: str = "") -> str:
    """One row per student per question with question_text, max_score and competency_code
    filled in; the instructor only types the scores. With no roster, a single block with a
    blank student_id is produced (blank rows are ignored on upload)."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(SCORE_SHEET_COLUMNS)
    for student in student_numbers or [""]:
        for q in questions:
            writer.writerow([
                student, q.question_text, "", f"{q.points:g}", submission_date,
                q.competency.competency_code if q.competency else "",
            ])
    return out.getvalue()
