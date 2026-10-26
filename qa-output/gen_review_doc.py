# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ---- Styles ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    return h

SEV_COLORS = {
    "HIGH": RGBColor(0xC0, 0x00, 0x00),
    "MEDIUM": RGBColor(0xB8, 0x6A, 0x00),
    "LOW": RGBColor(0x55, 0x55, 0x55),
}

def finding(num, title, severity, location, summary, scenario, fix):
    p = doc.add_heading(level=2)
    run = p.add_run(f"{num}. {title}")
    p2 = doc.add_paragraph()
    r = p2.add_run(f"Severity: {severity}")
    r.bold = True
    r.font.color.rgb = SEV_COLORS.get(severity, RGBColor(0,0,0))
    p2.add_run(f"      Location: {location}")
    for label, body in [("What", summary), ("Failure scenario", scenario), ("Suggested fix", fix)]:
        bp = doc.add_paragraph()
        br = bp.add_run(f"{label}: ")
        br.bold = True
        bp.add_run(body)

# ---- Title ----
title = doc.add_heading("AILAS — Code Review Findings", level=0)
sub = doc.add_paragraph()
sub.add_run("AI-Assisted Learning Analytics System  ·  Reviewer: Claude Code  ·  Date: 2026-06-03").italic = True

doc.add_heading("Scope & Method", level=1)
doc.add_paragraph(
    "The repository has no commit history yet (entire working tree untracked), so there was no diff to review. "
    "The whole codebase was treated as the review scope. Seven finder passes covered four high-risk domains "
    "(authentication/security, the AI risk/recommendation pipeline, FastAPI routers/data handling, and the React "
    "frontend). Each candidate was then re-verified by reading the actual source. Findings below are the verified, "
    "highest-severity issues, most severe first. Quality/maintainability notes follow the correctness/security bugs."
)

doc.add_heading("Summary", level=1)
doc.add_paragraph("Verified findings: 4 High, 5 Medium, 3 Low. The most important cluster is broken access control: "
    "several read endpoints authenticate the user but never authorize them, letting any logged-in account (including a "
    "student) read or modify data belonging to other sections/instructors. A second important issue is that LLM topic "
    "mapping silently never runs after a CSV upload because the background task reuses an already-closed DB session.")

# ============ HIGH ============
doc.add_heading("High Severity", level=1)

finding(1,
    "Broken access control: assessment TOS & approval-history readable by anyone",
    "HIGH",
    "backend/routers/assessments.py:264 (approval-history) and :278 (get_tos)",
    "Both endpoints depend only on get_current_user — there is no role check and no course-ownership check (_assert_course_access is used elsewhere but omitted here).",
    "A student, or an instructor from an unrelated college, calls GET /assessments/{any_uuid}/approval-history or /tos. The handler returns another section's reviewer feedback and Table-of-Specifications regardless of who owns the assessment — an authorization-bypass data leak (RA 10173 concern).",
    "Add an ownership/role guard: load the parent Course and call _assert_course_access(course, current_user), and restrict to instructor/PD/admin roles as appropriate.")

finding(2,
    "Broken access control: any grade submission readable by anyone",
    "HIGH",
    "backend/routers/grade_submissions.py:212 (get_submission) and :224 (approval-history)",
    "GET /grade-submissions/{id} and its approval-history depend only on get_current_user; no ownership or role scoping.",
    "A student or unrelated instructor enumerates/guesses a submission id and receives the full GradeSubmission (instructor_id, grading_period, notes) and approval history for a section they have no access to.",
    "Scope the query to the requesting instructor's submissions, or verify section/course access before returning.")

finding(3,
    "Broken access control: instructors can retag another instructor's assessment items",
    "HIGH",
    "backend/routers/assessments.py:490 (override_tag)",
    "override_tag requires only the instructor/admin role, with no check that the item belongs to a course the caller owns. It then bulk-updates every linked StudentScore.competency_id.",
    "Instructor B sends PUT /assessments/items/{item_id}/tag?competency_id=X for an item in Instructor A's course. B overwrites A's competency mapping and rewrites all of A's students' score-to-competency links, silently corrupting A's analytics and risk flags.",
    "Resolve the item's course and call _assert_course_access(course, current_user) before mutating; apply the same guard to the sibling /revert endpoint.")

finding(4,
    "Topic mapping silently never runs: background task reuses a closed DB session",
    "HIGH",
    "backend/routers/assessments.py:427 (background_tasks.add_task(_run_topic_mapping, assessment_id, db))",
    "The request-scoped Session from get_db() is passed into a BackgroundTask. FastAPI runs background tasks AFTER the response is sent, by which point get_db()'s finally-block has already called db.close(). Note auth.py's _bg_log correctly opens a fresh SessionLocal() instead — this path does not.",
    "Instructor uploads an assessment CSV. The endpoint returns 200, the session closes, then _run_topic_mapping queries/commits on the closed session and raises. The LLM competency mapping never completes, so every item stays unmapped, the mapping queue stays empty, and downstream risk/recommendation features have no competency data — with no error surfaced to the user.",
    "Open a new session inside the task (db = SessionLocal(); try/finally db.close()), mirroring _bg_log; pass only the assessment_id across the boundary.")

# ============ MEDIUM ============
doc.add_heading("Medium Severity", level=1)

finding(5,
    "CSV with max_score = 0 produces NaN normalized scores that poison analytics",
    "MEDIUM",
    "backend/services/assessment/csv_parser.py:25 and :33",
    "max_score is never validated as > 0. The score>max_score guard passes for 0>0 (False), then normalized_score = score/max_score*100 computes 0/0 = NaN, which is stored.",
    "An uploaded row with score=0,max_score=0 survives all validation and persists StudentScore.max_score=0 and a NaN normalized_score. Any later per-student/period average that includes the row becomes NaN or is skewed, distorting risk classification for that student.",
    "Reject rows where max_score <= 0 with a clear ValueError (e.g. add a check alongside the score>max_score guard).")

finding(6,
    "Login returns HTTP 500 instead of 401 on a malformed stored password hash",
    "MEDIUM",
    "backend/routers/auth.py:40 → backend/core/security.py:17 (verify_password)",
    "verify_password calls bcrypt.checkpw(plain, hashed.encode()), which raises ValueError('Invalid salt') if the stored hash is not a valid bcrypt string. The login flow only inspects the boolean result; the exception is uncaught.",
    "A user row seeded/imported with a non-bcrypt password_hash (empty string, legacy SHA-256 hex, etc.) crashes login with an unhandled 500 instead of a clean 401 — both a reliability bug and an account-existence oracle.",
    "Wrap verify_password in try/except and treat any exception as an authentication failure (return the generic 401).")

finding(7,
    "User-enumeration oracle: inactive account returns 403 after a correct password",
    "MEDIUM",
    "backend/routers/auth.py:44",
    "Wrong credentials return 401 'Invalid credentials'; an inactive account returns 403 'Account is inactive' — but only after the password has been verified as correct. The differing responses leak account state.",
    "An attacker submits a known email with the correct password for a disabled account and receives 403, confirming the email exists AND the password is valid (account merely deactivated), distinguishing it from the 401 wrong-password case.",
    "Return the same generic 401 for both invalid-credential and inactive-account cases (or fold the is_active check into the credential check before branching on a distinct status).")

finding(8,
    "Token revocation (logout) is not durable and not shared across workers",
    "MEDIUM",
    "backend/routers/auth.py:18 (_revoked_jtis) → backend/core/deps.py:25",
    "The JTI blacklist is an in-process Python set. It is lost on restart and not shared between Uvicorn/Gunicorn workers. (The inline comment acknowledges this as a single-session expedient.)",
    "A user logs out (jti added to the local set). A server restart, or simply the next request being routed to a different worker, clears/misses the entry, so the supposedly-revoked token is accepted again until its 30-minute exp. Logout provides no real revocation in any multi-worker or restarted deployment.",
    "Back the blacklist with the already-configured Redis (redis_url in config) keyed by jti with a TTL equal to the token's remaining lifetime; document the limitation if deferring.")

finding(9,
    "Recommender substring matching surfaces irrelevant resources",
    "MEDIUM",
    "backend/services/ai/recommender.py:70",
    "Tag matching uses bidirectional substring containment: any(term in tag or tag in term ...). Short competency codes/terms match unrelated tags as substrings.",
    "A weak competency with a short code/term such as 'co' matches resource tags like 'economics' or 'vocabulary' (term in tag), so unrelated videos/readings appear as recommendations with a nonzero relevance score on the student dashboard.",
    "Match on whole-token/exact equality (normalize and compare set membership), or require a minimum term length / word-boundary match instead of raw substring containment.")

# ============ LOW ============
doc.add_heading("Low Severity / Quality", level=1)

finding(10,
    "Frontend data race: useApi can apply a stale response over a newer one",
    "LOW",
    "frontend/src/hooks/useApi.js:13",
    "The fetch effect has no AbortController and no 'ignore' cleanup flag, so a late-resolving request for an old url can call setState after a newer request has already resolved.",
    "A student rapidly switches the course selector (A then B). If request A resolves after B, setData(A) overwrites B's data, so the dashboard shows course A's risk/scores while the selector reads course B.",
    "Add an ignore flag (or AbortController) in the effect and discard/abort the response in the cleanup function.")

finding(11,
    "At-Risk grade-decline trigger uses unweighted period means inconsistent with displayed averages",
    "LOW",
    "backend/services/ai/feature_engineering.py:135 (uses period_avgs from :124)",
    "grade_decline_pct is computed from simple unweighted per-period means (period_avgs), while the avg_score shown to students/instructors uses competency-mapped (or grade-weighted) averages. The two can disagree.",
    "In a course with grade weights, a prelim of one easy formative (100%) and a midterm of one hard summative (80%) yields a 20% simple decline → At-Risk fires, even though the weighted/competency average the user sees indicates On-Track — an unexplained mismatch between the flag and the dashboard numbers.",
    "Derive the decline from the same weighted/competency-based per-period figures used for display so the trigger and the visible averages are consistent.")

finding(12,
    "Fabricated per-period averages when a grading period has no data",
    "LOW",
    "backend/services/ai/feature_engineering.py:154-156",
    "Missing periods fall back to the previous period's value (prelim→avg_score, midterm→prelim, semi→midterm) instead of being reported as 'no data'.",
    "A student with scores only in the final period shows prelim_avg = midterm_avg = semi_final_avg = the overall average and a flat trend, so the dashboard presents fabricated prelim/midterm numbers for periods the student was never assessed in.",
    "Return None/omit periods that have no scores and let the UI render them as 'No data' rather than back-filling.")

doc.add_heading("Also worth addressing (noted, not detailed)", level=1)
for t in [
    "enrollments.py:_generate_student_number — read-then-write max()+1 can race under concurrent enrollment and collide with the year-prefixed seed, causing duplicate student_number (IntegrityError / shared id-hash).",
    "Rate limiter (core/limiter.py) keys on request.client.host with no X-Forwarded-For handling; behind the bundled nginx all clients collapse to the proxy IP, and state is per-process — so the login limit is both a shared-DoS risk and bypassable across workers.",
    "topic_mapper.py:json.loads has no None/parse guard of its own; it is safe only because the current caller wraps it in try/except — a trap for future callers.",
    "JWT 'role' claim is set to the UserRole enum object rather than its string value; harmless today (RBAC re-reads role from the DB) but makes the claim unreliable for any consumer that compares it to a literal string.",
    "StudentDashboard.jsx — auto-select-first-course effect and learning-gap rendering assume non-null numeric average_score and a stable courses identity; null/string values from the API can blank the dashboard.",
]:
    doc.add_paragraph(t, style="List Bullet")

out = os.path.join(os.environ["USERPROFILE"], "Desktop", "AILAS_Code_Review_Findings.docx")
doc.save(out)
print("SAVED:", out)
